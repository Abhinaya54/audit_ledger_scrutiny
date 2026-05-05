from docx import Document
from pathlib import Path
import re

path = Path('QA_TEST_REPORT.md')
doc = Document()
for line in path.read_text(encoding='utf-8').splitlines():
    if line.startswith('# '):
        doc.add_heading(line[2:].strip(), level=1)
    elif line.startswith('## '):
        doc.add_heading(line[3:].strip(), level=2)
    elif line.startswith('### '):
        doc.add_heading(line[4:].strip(), level=3)
    elif re.match(r'^\|.*\|$', line):
        doc.add_paragraph(line)
    elif line.startswith('---'):
        doc.add_page_break()
    elif line.strip() == '':
        doc.add_paragraph('')
    else:
        doc.add_paragraph(line)
output = Path('QA_TEST_REPORT.docx')
doc.save(output)
print(f'created {output.resolve()}')
